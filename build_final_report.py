"""Fill the supplied Word template using only direct paragraph/run formatting.

The document deliberately does not use Word Heading styles or Table Grid styles.
Heading outline levels are written directly so Word can build the table of contents.
"""
from copy import deepcopy
from pathlib import Path
from subprocess import run

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parent
TEMPLATE_DOC = ROOT / "docs" / "2026年计算机系统综合设计实验报告-雷军-2024302131025-尹玉文东-模板.doc"
TEMPLATE_DOCX = Path("/private/tmp/scpu_final_report_template.docx")
OUT = ROOT / "docs" / "雷军班_尹玉文东_2024302131025_实验报告.docx"
ARCH = Path("/private/tmp/scpu_final_arch.png")


def copy_rpr(target_run, source_run):
    if source_run._r.rPr is not None:
        target_run._r.get_or_add_rPr().getparent().replace(
            target_run._r.get_or_add_rPr(), deepcopy(source_run._r.rPr)
        )


def clone_format(target, source, outline=None):
    if source._p.pPr is not None:
        target._p.get_or_add_pPr().getparent().replace(
            target._p.get_or_add_pPr(), deepcopy(source._p.pPr)
        )
    if outline is not None:
        ppr = target._p.get_or_add_pPr()
        old = ppr.find(qn("w:outlineLvl"))
        if old is not None:
            ppr.remove(old)
        e = OxmlElement("w:outlineLvl")
        e.set(qn("w:val"), str(outline))
        ppr.append(e)


def set_text(p, text, source):
    p.clear()
    r = p.add_run(text)
    if source.runs:
        copy_rpr(r, source.runs[-1])
    return p


def add_text(doc, anchor, text, source, outline=None, align=None):
    p = doc.add_paragraph()
    p._p.getparent().remove(p._p)
    anchor._p.addprevious(p._p)
    clone_format(p, source, outline)
    if align is not None:
        p.alignment = align
    return set_text(p, text, source)


def add_page_break(doc, anchor):
    p = doc.add_paragraph()
    p._p.getparent().remove(p._p)
    anchor._p.addprevious(p._p)
    r = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r._r.append(br)
    return p


def add_field(doc, anchor, source):
    p = add_text(doc, anchor, "", source, align=WD_ALIGN_PARAGRAPH.LEFT)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    p._p.append(fld)
    return p


def set_cell(cell, text, source, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    clone_format(p, source)
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    if source.runs:
        copy_rpr(r, source.runs[-1])
    if bold:
        r.font.bold = True
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def border_table(table):
    tblpr = table._tbl.tblPr
    # Do not retain Word's built-in "Normal Table" style; all table geometry
    # and borders are set directly to follow the supplied template convention.
    old_style = tblpr.find(qn("w:tblStyle"))
    if old_style is not None:
        tblpr.remove(old_style)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        borders.append(e)
    tblpr.append(borders)


def add_table(doc, anchor, headers, rows, source, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t._tbl.getparent().remove(t._tbl)
    anchor._p.addprevious(t._tbl)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    border_table(t)
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.width = Cm(widths[j])
        set_cell(cell, h, source, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = t.add_row().cells
        for j, value in enumerate(row):
            cells[j].width = Cm(widths[j])
            set_cell(cells[j], value, source, align=WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    return t


def add_picture(doc, anchor, path, width_cm):
    p = doc.add_paragraph()
    p._p.getparent().remove(p._p)
    anchor._p.addprevious(p._p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    return p


def add_code(doc, anchor, code, source):
    """Insert a directly formatted code paragraph; never inherit body justify."""
    p = add_text(doc, anchor, code, source)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.left_indent = Cm(0.74)
    pf.right_indent = Cm(0.25)
    pf.space_before = Pt(3)
    pf.space_after = Pt(5)
    pf.line_spacing = 1.0
    for r in p.runs:
        r.font.name = "Menlo"
        r._r.rPr.rFonts.set(qn("w:ascii"), "Menlo")
        r._r.rPr.rFonts.set(qn("w:hAnsi"), "Menlo")
        r.font.size = Pt(8.5)
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    ppr.append(shd)
    return p


def arch_image():
    dot = r'''digraph G {
      graph [rankdir=LR, bgcolor="white", pad="0.08", nodesep="0.26", ranksep="0.36"];
      node [shape=box, style="rounded,filled", fontname="PingFang SC", fontsize=11, color="#444444", fillcolor="#F4F4F4", margin="0.10,0.06"];
      edge [fontname="PingFang SC", fontsize=8, color="#555555", arrowsize=0.65];
      imem [label="指令 ROM\nCOE"];
      cpu [label="五级流水线 CPU\nIF · ID · EX · MEM · WB\n前推/冒险/中断异常", fillcolor="#D9EAD3"];
      mio [label="MIO 总线\n地址译码", fillcolor="#D9EAF7"];
      dmem [label="数据 RAM"];
      vga [label="VGA 文本显示\n80×60"];
      ps2 [label="PS/2 接口"];
      timer [label="游戏定时器\nINT"];
      gpio [label="LED/数码管"];
      app [label="Dino 应用\n帧更新/状态机", fillcolor="#FCE5CD"];
      imem -> cpu [label="inst"]; cpu -> dmem [label="访存"]; cpu -> mio [label="MMIO"];
      ps2 -> mio; mio -> vga; mio -> gpio; timer -> cpu [label="IRQ"];
      cpu -> app [style=dashed, label="执行"];
    }'''
    src = Path("/private/tmp/scpu_final_arch.dot")
    src.write_text(dot, encoding="utf-8")
    run(["dot", "-Tpng", "-Gdpi=180", str(src), "-o", str(ARCH)], check=True)


def main():
    run(["textutil", "-convert", "docx", "-output", str(TEMPLATE_DOCX), str(TEMPLATE_DOC)], check=True)
    doc = Document(TEMPLATE_DOCX)
    ps = doc.paragraphs
    by = {p.text: p for p in ps if p.text}
    # The user may have typed a draft abstract into the source template before
    # rerunning the builder.  Use the original body paragraph when present,
    # otherwise use the first abstract body paragraph solely as a formatting
    # reference; all generated text is inserted into the separate output file.
    body = by.get("本实验……")
    if body is None:
        body = next(p for p in ps if p.text.startswith("本实验的实验目的是"))
    h1 = by["3 系统架构设计"]
    h2 = by["1.1 实验目的"]
    cap = by["信号名"]
    abstract_title = by["摘  要"]
    toc_title = by["目  录"]
    reference = by["参考文献"]

    # Cover information: preserve the template’s direct formatting.
    for p in ps:
        if p.text == "专 业 名 称   ：":
            set_text(p, "专 业 名 称   ：计算机科学与技术(雷军班)", p)
        elif p.text == "学 生 学 号   ：":
            set_text(p, "学 生 学 号   ：2024302131025", p)
        elif p.text == "学 生 姓 名   ：":
            set_text(p, "学 生 姓 名   ：尹玉文东", p)

    # Remove all original guidance text from abstract through references.
    start = abstract_title._p
    end = reference._p
    children = list(start.getparent())
    a, b = children.index(start), children.index(end)
    for e in children[a:b]:
        e.getparent().remove(e)

    # Re-find the reference anchor after deleting the old body.
    reference = next(p for p in doc.paragraphs if p.text == "参考文献")
    arch_image()
    zoom = ROOT / "docs" / "interrupt_timer_zoom.png"
    ret = ROOT / "docs" / "interrupt_timer_return.png"

    # Abstract and TOC.
    add_text(doc, reference, "摘  要", abstract_title, outline=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, reference, "本实验完成了一套基于 RISC-V RV32I 指令集的 FPGA SoC 系统，并在 Nexys A7 开发板上运行 Dino 小游戏。系统以五级流水线 CPU 为核心，围绕“能正确执行程序、能访问外设、能观察并定位错误”三个目标组织设计。硬件部分包括指令 ROM、数据 RAM、寄存器文件、算术逻辑单元、流水线寄存器、MIO 总线、PS/2 键盘接口、VGA 文本显示接口、LED/数码管接口和游戏定时器；软件部分包括裸机启动代码、异常和中断入口以及 Dino 应用程序。", body)
    add_text(doc, reference, "处理器包含取指、译码、执行、访存和写回五个阶段。针对寄存器数据相关，设计以前推单元选择 EX/MEM 或 MEM/WB 中的最新数据；针对无法直接前推的 load-use 情形，由冒险检测单元暂停前级并插入气泡。对控制相关和 trap，系统在 EX 阶段集中处理 PC 重定向和流水线冲刷，避免错误路径指令产生写寄存器、写存储器或写外设等副作用。", body)
    add_text(doc, reference, "在异常和中断方面，系统支持非法指令异常、ECALL 异常和计时器中断。计时器产生请求后，CPU 锁存 pending 位，并在 INTMASK 允许且不处于 trap 状态时跳转至 0x00000340；进入处理程序时保存 SEPC 与 SCAUSE，处理程序更新 frame_ticks，主程序据此进行游戏状态更新和局部重绘。仿真波形表明中断响应时 SCAUSE 为 06、INTMASK 为 40，PC 能够正确转入计时器向量并在返回后恢复执行。", body)
    add_text(doc, reference, "关键词：RISC-V；五级流水线；SoC；中断；VGA；PS/2；Nexys A7", body)
    add_page_break(doc, reference)
    add_text(doc, reference, "目  录", toc_title, outline=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_field(doc, reference, body)
    add_page_break(doc, reference)

    # 1 Introduction.
    add_text(doc, reference, "1 引言", h1, outline=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, reference, "1.1 实验目的", h2, outline=1)
    add_text(doc, reference, "本实验的首要目标是把处理器原理课中相对独立的数据通路、控制器、流水线相关、存储器访问和异常处理知识落实为一个可综合、可下载、可运行的系统。单独观察 ALU、寄存器堆或控制器时，往往只能验证某一条指令的局部功能；在 SoC 中，这些模块必须经由流水线寄存器、时钟和控制信号协作，才能保证一段真实程序的执行结果正确。", body)
    add_text(doc, reference, "第二个目标是完成 RV32I 基础指令的硬件支持，并验证算术逻辑、立即数、分支跳转以及字节、半字、字访存等不同路径。实验并不把“能够运行一条指令”作为终点，而是要求处理器在连续指令、存在数据依赖、分支改变控制流和访问外设时仍保持正确。因此，前推、停顿、冲刷和 valid 控制是本项目的重要内容。", body)
    add_text(doc, reference, "第三个目标是建立软硬件协同流程。应用程序经交叉编译和链接后生成 COE 初始化文件，指令 ROM 在上电后取出相应机器码；软件通过 load/store 访问统一的内存映射地址，硬件再将访问分发给 RAM、VGA、PS/2、LED、数码管和定时器。这样能够把软件现象与硬件信号建立对应关系，既便于调试，也符合 SoC 系统设计的基本工作方式。", body)
    add_text(doc, reference, "最后，本实验以 Dino 游戏作为综合应用。游戏并非只用于展示，而是同时覆盖了循环、条件分支、状态机、键盘输入、屏幕输出、计时器中断和分数显示等场景。游戏能够运行说明 CPU、MMIO、外设和应用程序之间的接口已完成闭环；出现异常时，也可以借助波形、显示信息和模块化测试逐级缩小故障范围。", body)
    add_text(doc, reference, "1.2 国内外研究现状", h2, outline=1)
    add_text(doc, reference, "RISC-V 是开放指令集规范，其基础整数指令集规模适中、编码规则清晰，适合用于处理器教学和原型验证。与直接使用封闭体系结构的现成软核相比，围绕 RV32I 从数据通路到应用程序完成一次实现，可以使学习者直观看到一条 C 语句或汇编指令如何被编译、取指、译码、执行、访存和写回。开放规范也使课程能够将重点放在微体系结构本身，而不是特定商业工具链的黑盒实现上。", body)
    add_text(doc, reference, "在国外的体系结构教学中，常见路径是先用软件模拟器或单周期数据通路讲解 ISA，再逐步加入流水线、前推、停顿、缓存、异常和中断。实际硬件实验通常会选择 FPGA 作为验证平台，将抽象的寄存器值、PC 和控制信号与可观察的显示、按键和串行输出联系起来。这样既可以用仿真观察周期级细节，也可以在实板上验证时钟、引脚和外设接口。", body)
    add_text(doc, reference, "国内相关课程同样强调“从 CPU 到 SoC”的连续性。单周期和流水线 CPU 解决的是指令执行问题，而内存映射 I/O、定时器、显示器和键盘接口解决的是程序如何与外部世界交互的问题。对于教学项目而言，文本模式 VGA 比图形帧缓存更容易控制资源与调试范围；PS/2 扫描码和数码管显示又提供了直观的输入输出验证手段。", body)
    add_text(doc, reference, "本项目采用小规模裸机 SoC 的定位，不涉及操作系统、虚拟内存和复杂中断控制器，而把重点放在五级流水线、异常/中断优先级、外设地址译码和应用集成上。这一取舍有利于在有限课时内完成从 RTL 到 FPGA 的完整闭环，也使每一项测试结论都能回溯到具体模块、具体地址或具体波形。", body)

    # 2 Environment.
    add_text(doc, reference, "2 实验环境介绍", h1, outline=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, reference, "2.1 Verilog HDL", h2, outline=1)
    add_text(doc, reference, "Verilog HDL 用于描述 CPU 数据通路、控制器、MIO 总线、PS/2 接口、VGA 文本显示和游戏定时器，并用于编写顶层仿真测试平台。项目按功能划分 RTL 文件：CPU 内部包括 PC、寄存器文件、立即数扩展、ALU、控制器、流水线寄存器、前推单元、冒险检测和异常中断单元；板级部分负责时钟、复位、ROM、RAM、VGA、PS/2 和外设连接。", body)
    add_text(doc, reference, "在编码过程中，组合逻辑主要使用连续赋值或 always @(*) 描述，时序状态使用时钟触发的 always 块描述。对于流水线设计，除数据和控制信号外，还需要把 valid 位随流水寄存器传递；这样在插入气泡或发生冲刷时，即使个别控制字段保留旧值，也不会对寄存器、RAM 或 MMIO 产生实际写操作。", body)
    add_text(doc, reference, "2.2 交叉编译与仿真工具", h2, outline=1)
    add_text(doc, reference, "游戏程序使用 C 语言和少量汇编代码编写。启动汇编负责设置栈指针、清零 BSS 区域并进入 main；异常和中断入口保存必要现场、更新状态并执行返回指令。C 语言程序负责键盘扫描码解释、Dino 状态机、障碍移动、碰撞判断、得分处理以及 VGA 文本单元写入。", body)
    add_text(doc, reference, "工程在本机使用 Homebrew LLVM/LLD 工具链：clang 的目标为 riscv32-unknown-elf，链接器为 ld.lld，编译选项指定 -march=rv32i 和 -mabi=ilp32。链接脚本将启动代码、普通代码和 trap 入口布置到指令 ROM 对应区域，构建脚本再把裸机 ELF 转换为二进制和 COE。生成 COE 后，仿真与 Vivado 工程使用相同的指令初始化内容，从而减少软件版本不一致造成的问题。", body)
    add_text(doc, reference, "顶层仿真采用 Icarus Verilog。测试脚本负责读取 COE、初始化 ROM/RAM、施加复位、模拟键盘扫描码和设置检查条件；运行时可选输出 VCD。GTKWave 用于查看 PC、SEPC、SCAUSE、STATUS、INTMASK、pending 位和中断控制信号，特别适合确认中断请求、向量跳转、处理程序执行和返回之间的时序关系。", body)
    add_text(doc, reference, "2.3 Vivado 2018.3", h2, outline=1)
    add_text(doc, reference, "Vivado 2018.3 用于综合、实现、生成 bitstream 和下载 FPGA。工程中需要把 RTL、ROM 初始化文件和 XDC 约束统一纳入工程，尤其要保证顶层端口名称与 Nexys A7 引脚约束一致。综合和实现完成后，可检查时钟约束、资源使用和关键路径，再生成 bitstream 下载到 FPGA。", body)
    add_text(doc, reference, "Vivado 的作用不仅是生成配置文件，也用于检查板级连接问题。例如，VGA 的 RGB 与行场同步信号必须使用正确的引脚和电平标准；PS/2 接口是双向开漏形式，需要与顶层三态连接方式匹配。实板出现显示不稳定、输入无响应或程序不运行时，应先检查 bitstream、时钟和约束，再回到仿真定位 RTL 问题。", body)
    add_text(doc, reference, "2.4 Nexys A7", h2, outline=1)
    add_text(doc, reference, "Nexys A7 是本实验的目标开发板。系统使用其板载时钟、LED、七段数码管、VGA 和 PS/2 接口；VGA 采用 80×60 文本单元显示游戏状态和得分。板级顶层首先对时钟和复位进行整理，再向 CPU、ROM、RAM 和外设分发所需时钟。游戏定时器默认关闭，软件写入控制地址后才开始产生稳定的帧节拍，避免普通测试程序被意外中断。", body)
    add_text(doc, reference, "在应用验证中，VGA 用于观察 READY、运行、GAME OVER 和得分等可见结果；PS/2 键盘用于触发跳跃、重启等操作；LED 与七段数码管用于输出状态或得分辅助信息。多种输出路径同时存在的好处是：即使 VGA 画面尚未完全正确，也可以通过 LED、数码管和仿真日志判断程序是否在继续执行。", body)
    add_text(doc, reference, "2.5 软件构建与板级验证流程", h2, outline=1)
    add_text(doc, reference, "本项目的验证流程分为四步。第一步，在模块级或 CPU 级仿真中验证基础指令、访存和流水线冒险；第二步，编译应用程序并生成 COE，使用顶层测试平台检查 MMIO 写入、PS/2 输入和中断波形；第三步，在 Vivado 中综合、实现并生成 bitstream；第四步，将 bitstream 下载至 Nexys A7，连接 VGA 显示器和 PS/2 键盘，观察游戏运行结果。每一步都保留可重复的输入文件和测试条件，便于在出现问题时回退定位。", body)
    add_table(doc, reference, ["类别", "环境", "用途"], [
        ("硬件平台", "Nexys A7", "下载 SoC，连接 VGA 和 PS/2 键盘"),
        ("FPGA 工具", "Vivado 2018.3", "综合、实现、下载和时序检查"),
        ("交叉编译器", "clang（riscv32-unknown-elf）", "生成 RV32I 应用和 COE"),
        ("链接器", "ld.lld", "裸机程序链接"),
        ("仿真工具", "Icarus Verilog、GTKWave", "顶层仿真与 VCD 波形分析"),
    ], cap, [3.1, 5.4, 6.6])
    add_text(doc, reference, "表 2-1 实验环境", cap, align=WD_ALIGN_PARAGRAPH.CENTER)

    # 3 Architecture.
    add_text(doc, reference, "3 系统架构设计", h1, outline=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, reference, "3.1 总体设计", h2, outline=1)
    add_text(doc, reference, "最终系统为五级流水线 RISC-V SoC。IF 阶段从指令 ROM 取指；ID 阶段完成寄存器读取、立即数扩展和控制信号生成；EX 阶段完成 ALU 运算、分支判断以及 trap 判定；MEM 阶段访问数据 RAM 或内存映射外设；WB 阶段将运算结果、存储器读数或 PC+4 写回寄存器堆。前推单元从 EX/MEM、MEM/WB 选择最新操作数，冒险检测单元对 load-use 相关暂停前级并插入气泡。", body)
    add_picture(doc, reference, ARCH, 15.2)
    add_text(doc, reference, "图 3-1 SoC 总体架构", cap, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, reference, "3.2 地址映射与向量入口", h2, outline=1)
    add_table(doc, reference, ["地址/入口", "功能模块", "说明"], [
        ("0x0000_0000 起", "指令 ROM", "保存 RV32I 程序和中断入口代码"),
        ("0xC000_0000", "VGA 文本显存", "按 row×80+col 写入字符和属性"),
        ("0xD000_0000", "PS/2 键盘", "读取扫描码和就绪状态"),
        ("0xE000_0000", "GPIO/七段数码管", "读取开关按键并显示得分"),
        ("0xF000_0000", "LED", "显示游戏状态和辅助结果"),
        ("0xFFFF_FE00", "游戏定时器", "使能或关闭帧节拍"),
        ("0xFFFF_FF00", "INTMASK", "设置计时中断允许位"),
        ("0x0000_0300", "非法指令异常", "SCAUSE=01"),
        ("0x0000_0320", "ECALL 异常", "SCAUSE=02"),
        ("0x0000_0340", "计时器中断", "SCAUSE=06，进入帧中断处理程序"),
    ], cap, [3.1, 4.0, 8.0])
    add_text(doc, reference, "表 3-1 地址映射与 trap 向量", cap, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, reference, "3.3 功能模块及接口", h2, outline=1)
    add_table(doc, reference, ["模块", "功能描述", "主要输入", "主要输出"], [
        ("PC/NPC", "保存 PC 并选择顺序、分支、跳转或 trap 目标。", "clk、reset、redirect、trap_vector", "PC_out、PC+4"),
        ("RF", "保存 32 个通用寄存器，x0 恒为 0。", "rs1、rs2、rd、wd、RegWrite", "rs1_data、rs2_data"),
        ("ALU", "执行算术、逻辑、比较和地址计算。", "ALU_A、ALU_B、ALUOp", "alu_result、zero"),
        ("控制器", "识别指令并产生流水线控制信号。", "inst", "ALUOp、RegWrite、MemWrite"),
        ("前推单元", "选择最新操作数，解决可前推 RAW 相关。", "流水寄存器号、RegWrite", "forwardA、forwardB"),
        ("冒险检测", "检测 load-use 相关并插入气泡。", "rs、rd、MemRead", "stall、flush"),
        ("DM/MIO", "处理不同宽度访存并完成地址译码。", "addr、data、mem_w、dm_ctrl", "Data_read、外设写使能"),
        ("异常中断控制", "保存 trap 状态、选择入口并冲刷流水线。", "EX cause、INT、STATUS、INTMASK", "SEPC、SCAUSE、trap_vector"),
        ("PS/2", "接收键盘扫描码。", "PS2C、PS2D、read", "key、ready、scan_log"),
        ("VGA", "把文本显存转换为像素与同步信号。", "MMIO 写入、时钟", "VGA_R/G/B、HS、VS"),
        ("游戏定时器", "产生固定周期的帧中断。", "clk、enable", "tick_irq"),
    ], cap, [2.2, 5.1, 4.2, 4.0])
    add_text(doc, reference, "表 3-2 功能模块、功能描述与主要接口", cap, align=WD_ALIGN_PARAGRAPH.CENTER)

    # The teacher requires the important functional units to be documented one
    # by one, rather than only listed in a summary table.
    module_details = [
        ("3.4 PC、NPC 与取指模块", "PC 寄存器在时钟沿保存下一条待取指令的地址。复位后 PC 指向指令 ROM 起始地址；无控制转移时 NPC 为 PC+4；分支成立、JAL/JALR、异常入口、计时器中断和 ERET/ERETN 分别提供不同的重定向目标。为了保证流水线状态一致，发生 trap 或错误路径跳转时，PC 更新与 IF/ID、ID/EX 的冲刷在同一控制分支中完成。",
         [("clk", "输入", "CPU 时钟"), ("reset", "输入", "复位信号"), ("redirect_pc", "输入", "分支或跳转目标地址"), ("trap_vector", "输入", "异常/中断入口地址"), ("PC_out", "输出", "送往指令 ROM 的当前 PC"), ("pc_plus4", "输出", "顺序下一条指令地址")]),
        ("3.5 寄存器文件 RF", "寄存器文件包含 32 个 32 位通用寄存器，提供两个组合读端口和一个同步写端口。读地址由 rs1、rs2 给出，写地址由 rd 给出；当 RegWrite 有效且 rd 不为 0 时，在写回阶段写入写回数据。对 x0 的写入被忽略，读 x0 始终返回 0。译码阶段读取的数据会进入 ID/EX 寄存器，并在 EX 阶段可被前推结果覆盖。",
         [("rs1/rs2", "输入", "两个源寄存器编号"), ("rd", "输入", "写回目的寄存器编号"), ("wd", "输入", "写回数据"), ("RegWrite", "输入", "寄存器写使能"), ("rs1_data", "输出", "源操作数 1"), ("rs2_data", "输出", "源操作数 2")]),
        ("3.6 ALU、立即数扩展与主控制器", "主控制器根据 opcode、funct3 和 funct7 生成 ALUOp、ALUSrc、RegWrite、MemRead、MemWrite、WDSel、Branch、JAL/JALR 等控制信号。立即数扩展模块分别支持 I、S、B、U、J 型立即数，并按指令格式进行符号扩展或拼接。ALU 既用于算术逻辑计算，也用于访存地址计算和分支比较；比较结果配合分支控制信号产生真实的 PC 重定向条件。",
         [("inst", "输入", "当前译码指令"), ("rs1_data/rs2_data", "输入", "寄存器操作数"), ("imm", "输入", "扩展后的立即数"), ("ALUOp", "输入", "ALU 运算选择"), ("alu_result", "输出", "算术或地址计算结果"), ("zero/compare", "输出", "分支比较结果")]),
        ("3.7 流水线寄存器、前推与冒险检测", "系统使用 IF/ID、ID/EX、EX/MEM 和 MEM/WB 四组流水寄存器保存指令相关数据和控制信号，并用 valid 位标识该级是否为有效指令。前推单元检测当前 EX 阶段的 rs1、rs2 是否命中后续级即将写回的 rd；冒险检测单元专门处理 load-use 情形。分支、JAL/JALR、异常和中断的重定向会将错误路径寄存器置为 bubble，防止错误路径进入 MEM/WB 产生外部副作用。",
         [("id_ex_rs1/rs2", "输入", "EX 阶段源寄存器号"), ("ex_mem_rd", "输入", "EX/MEM 目的寄存器号"), ("mem_wb_rd", "输入", "MEM/WB 目的寄存器号"), ("MemRead", "输入", "load 指令标识"), ("forwardA/forwardB", "输出", "ALU 操作数选择"), ("stall/flush", "输出", "停顿和气泡控制")]),
        ("3.8 数据存储器与 MIO 总线", "MEM 阶段依据地址高位把访问分为数据 RAM 和外设访问。DM 控制器完成 lb、lbu、lh、lhu、lw 的读数据选择及扩展，也完成 sb、sh、sw 的写数据对齐。MIO 总线对不同地址段产生外设片选和读写使能，使 CPU 以统一的 load/store 指令访问 VGA、PS/2、LED、数码管、INTMASK 与游戏定时器。",
         [("Addr_out", "输入", "MEM 阶段访问地址"), ("Data_out", "输入", "MEM 阶段写数据"), ("mem_w", "输入", "写请求"), ("dm_ctrl", "输入", "访问宽度及扩展方式"), ("Data_read", "输出", "送回 WB 的读数据"), ("peripheral_we", "输出", "外设写使能")]),
        ("3.9 单级异常与中断控制模块", "异常中断控制模块维护 SEPC、SCAUSE、STATUS、INTMASK 和 int_pending。INT 输入到来时先锁存为 pending，避免短脉冲丢失；当 INTMASK[6] 打开且不在 trap 状态时，ExceptionUnit 输出 SCAUSE=06 和 0x00000340。非法指令和 ECALL 分别输出 0x00000300、0x00000320，并优先于异步中断。进入 trap 后 STATUS[0] 置 1，从而禁止嵌套响应，构成单级模型。",
         [("EX_SCAUSE", "输入", "EX 阶段同步异常原因"), ("INT", "输入", "计时器中断请求"), ("INTMASK", "输入/状态", "中断使能位"), ("STATUS", "输入/状态", "EXL 状态位"), ("SEPC", "输出/状态", "保存的返回地址"), ("SCAUSE", "输出/状态", "保存的 trap 原因"), ("trap_vector", "输出", "异常或中断入口")]),
        ("3.10 PS/2、VGA 和游戏定时器", "PS/2 接口把键盘串行数据还原为扫描码，并通过 MMIO 提供 ready、当前键值和扫描记录。VGA 模块把 CPU 写入的文本单元缓存转换为行场同步与 RGB 输出，游戏程序以 80×60 网格坐标写入标题、地面、角色、障碍和得分。游戏定时器在软件写入使能后按固定周期生成单周期 tick_irq，CPU 接收后转化为计时器中断。",
         [("PS2C/PS2D", "输入", "PS/2 串行时钟与数据"), ("MMIO_VGA", "输入", "文本显存写地址和数据"), ("enable_we/enable_data", "输入", "定时器控制写入"), ("tick_irq", "输出", "计时器中断脉冲"), ("VGA_R/G/B", "输出", "VGA 颜色输出"), ("VGA_HS/VGA_VS", "输出", "VGA 同步信号")]),
    ]
    for title, desc, ports in module_details:
        add_text(doc, reference, title, h2, outline=1)
        add_text(doc, reference, desc, body)
        add_table(doc, reference, ["信号名", "方向", "说明"], ports, cap, [3.6, 2.0, 9.7])
        add_text(doc, reference, "表 " + title.split(" ")[0] + " 模块接口", cap, align=WD_ALIGN_PARAGRAPH.CENTER)

    # 4 Critical modules.
    add_text(doc, reference, "4 关键模块设计", h1, outline=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    for title, text, code in [
        ("4.1 DM 读写控制器", "DM 控制器依据 dm_ctrl 和地址低位完成字节、半字和字的对齐访问。读操作先从 32 位数据中选择目标字节或半字，再按 lb/lbu、lh/lhu 的语义进行符号扩展或零扩展；写操作按照 sb、sh、sw 的宽度生成相应写数据。", "// 根据地址低两位选择当前访问的字节\ncase (addr[1:0])\n  2'b00: byte_data = Data_in[7:0];\n  2'b01: byte_data = Data_in[15:8];\n  2'b10: byte_data = Data_in[23:16];\n  default: byte_data = Data_in[31:24];\nendcase\n// unsigned_load 为 1 时执行零扩展，否则按符号位扩展\nread_data = unsigned_load ? {24'b0, byte_data} :\n                           {{24{byte_data[7]}}, byte_data};"),
        ("4.2 前推单元", "前推单元比较 ID/EX 源寄存器与 EX/MEM、MEM/WB 目的寄存器。若 EX/MEM 已产生可写回结果，则优先选择该结果；否则在 MEM/WB 命中时选择写回数据。x0 不参与前推。", "// EX/MEM 的结果更新更晚，因此具有更高前推优先级\nif (ex_mem_regwrite && ex_mem_rd != 0 && ex_mem_rd == id_ex_rs1)\n    forwardA = 2'b10;\n// 未命中 EX/MEM 时，再检查 MEM/WB 写回结果\nelse if (mem_wb_regwrite && mem_wb_rd != 0 && mem_wb_rd == id_ex_rs1)\n    forwardA = 2'b01;\n// 两级均不相关，直接使用 ID/EX 保存的寄存器值\nelse\n    forwardA = 2'b00;"),
        ("4.3 冒险检测", "load 指令的读数据在 MEM 阶段末才可获得，紧随其后的使用者无法仅靠前推取得正确值。冒险检测单元因此保持 PC 和 IF/ID 寄存器，并向 ID/EX 注入气泡。", "// EX 阶段为 load，且目的寄存器被 ID 阶段指令使用\nload_use = id_ex_memread && (id_ex_rd != 0) &&\n           ((id_ex_rd == if_id_rs1) || (id_ex_rd == if_id_rs2));\nif (load_use) begin\n    // 保持取指和译码内容不变，等待数据从 MEM 阶段返回\n    pc_write = 1'b0;\n    if_id_write = 1'b0;\n    // 向 EX 级注入无副作用的气泡\n    id_ex_flush = 1'b1;\nend"),
        ("4.4 单级中断与异常控制器", "控制器在 EX 阶段统一处理 trap。非法指令和 ECALL 是同步异常，优先于异步计时中断。进入 trap 时保存 SEPC 和 SCAUSE，置位 STATUS[0]，清空错误路径和当前 EX 指令的副作用；ERET 或 ERETN 返回后清除 STATUS[0]。", "// STATUS[0] 为 1 表示正在处理 trap，禁止嵌套响应\nif (!STATUS[0]) begin\n  // 同步异常优先：非法指令或 ECALL 直接选择各自入口\n  if (EX_SCAUSE != SCAUSE_NONE) begin\n      trap_vector = exception_vector;\n      trap_set = 1'b1;\n  // 同步异常不存在时，检查已锁存且已使能的计时器请求\n  end else if (INT_PEND[6] && INTMASK[6]) begin\n      trap_vector = 32'h0000_0340;\n      trap_set = 1'b1;\n      int_signal = 1'b1;\n  end\nend"),
        ("4.5 PS/2、VGA 与游戏帧更新", "PS/2 接口由尹玉文东负责，向软件提供扫描码、就绪状态和扫描记录；VGA 接口由钱瑞恒负责，将文本显存映射为显示输出。两人共同设计游戏状态机和程序逻辑。计时器中断只发布 frame_ticks，主循环在获得节拍后更新恐龙、障碍、碰撞与得分，并只重绘改变的文本单元。", "// frame_ticks 由计时器中断服务程序递增\nif (frame_ticks != 0u) begin\n    // 合并多个延迟到达的节拍，避免一次循环重复重绘\n    frame_ticks = 0u;\n    // 更新跳跃速度、障碍坐标和碰撞状态\n    update_physics();\n    update_obstacle();\n    // 仅写入发生变化的 VGA 文本单元\n    render_changed_cells();\nend"),
    ]:
        add_text(doc, reference, title, h2, outline=1)
        add_text(doc, reference, text, body)
        add_code(doc, reference, code, body)

    # 5 Tests.
    add_text(doc, reference, "5 测试及结果分析", h1, outline=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, reference, "5.1 测试代码及分析", h2, outline=1)
    add_text(doc, reference, "测试按照“基础 CPU 功能—板级外设接口—trap 时序—应用程序”四个层次进行。每一层都给出对应测试文件、输入刺激、观察信号和通过判据。这样可以避免直接运行复杂游戏时，无法判断故障来自 CPU、MMIO、外设还是应用逻辑的问题。统一测试程序用于验证课程规定的基础指令和流水线行为；由本组编写的顶层测试平台和 Dino 程序用于验证中断、PS/2、VGA 与应用集成。", body)
    add_text(doc, reference, "5.1.1 基础指令与流水线测试", h2, outline=2)
    add_text(doc, reference, "基础 CPU 测试使用 sim/sccomp_tb.v 配合课程测试 COE 执行。该测试平台向 CPU 提供时钟、复位、指令存储器和数据存储器模型，并在程序运行结束时观察寄存器写回、RAM 写入及显示标记。测试覆盖算术逻辑指令、立即数指令、分支跳转、JAL/JALR、字节/半字/字访存，以及连续指令之间的 RAW 相关。对前推有效的相关，下一条指令不应插入多余停顿；对 load-use 相关，PC 和 IF/ID 应保持一个周期，随后使用正确的 load 数据继续执行。", body)
    add_code(doc, reference, "// sim/sccomp_tb.v：导出 CPU 级波形，便于检查流水线寄存器\ninitial begin\n    $dumpfile(\"build/sccomp_tb.vcd\");\n    $dumpvars(0, sccomp_tb);\n    // 复位后由测试 COE 驱动指令序列，直到测试程序结束\nend", body)
    add_text(doc, reference, "通过判据为：测试程序给出的阶段标记依次出现；寄存器 x0 始终为 0；分支后的错误路径不会写回；不同宽度的读写数据与小端存储规则一致。若该层失败，应先检查指令译码、立即数、ALU 控制和写回选择，再检查前推与冒险控制。", body)
    add_text(doc, reference, "5.1.2 顶层 SoC 与 MMIO 测试", h2, outline=2)
    add_text(doc, reference, "顶层测试使用 sim/top_board_tb.v 和 sim/run_top_board_sim.py。脚本先解析 COE 文件中的 memory_initialization_vector，将其转换为仿真 ROM 可读取的文本文件；随后调用 Icarus Verilog 编译 board/top.v 及相关外设模型。top_board_tb.v 监视 CPU 写入 0xE0000000 的数值、VGA 文本显存写地址、PC 和外设状态，因此可以在不依赖实板显示器的情况下验证 MIO 地址译码和外设写入。", body)
    add_code(doc, reference, "# 运行 Dino 顶层仿真：检查定时器使能、0x340 向量和得分变化\npython3 sim/run_top_board_sim.py \\\n  --imem coe/board/I_dino_game.coe \\\n  --sw 0000 --max-cycles 1000000 \\\n  --timer-period 20000 --check-dino-tick --dump-vcd\n# --dump-vcd 生成 build/top_board_tb.vcd，供 GTKWave 观察", body)
    add_text(doc, reference, "顶层测试的通过条件不是只看程序是否停止，而是同时满足三项：软件向游戏定时器控制地址写入使能位；PC 至少一次进入 0x00000340；游戏得分寄存器或显示 MMIO 出现首个递增结果。三项同时成立说明应用已经完成初始化，定时器请求进入 CPU，处理程序返回后主循环继续推进一帧。", body)
    add_text(doc, reference, "5.1.3 PS/2 键盘输入测试", h2, outline=2)
    add_text(doc, reference, "PS/2 测试仍在 top_board_tb.v 中完成。测试平台通过 send_ps2_byte 任务产生起始位、8 位数据、奇偶校验位和停止位，从而模拟键盘发送的扫描码。发送空格键扫描码 0x29 后，PS/2 接口应使 ready 有效，CPU 读取 0xD0000000 后获得键值并确认该事件；Dino 程序则把该输入解释为跳跃动作。对于断码和扩展码，软件使用 break_pending 与 ext_pending 记录前缀，避免把释放事件或方向键前缀误识别为普通按键。", body)
    add_code(doc, reference, "// top_board_tb.v：向 PS/2 接口按串行格式发送一个扫描码\ntask automatic send_ps2_byte;\n    input [7:0] code;\n    begin\n        // 先发起始位，再按低位在前依次发送 8 位扫描码\n        send_ps2_bit(1'b0);\n        for (bit_idx = 0; bit_idx < 8; bit_idx = bit_idx + 1)\n            send_ps2_bit(code[bit_idx]);\n        // 补充奇偶校验位和停止位\n        send_ps2_bit(parity_bit);\n        send_ps2_bit(1'b1);\n    end\nendtask", body)
    add_text(doc, reference, "PS/2 接口的通过判据为：仿真输出中出现被发送的扫描码，MMIO 返回值同时包含 ready 状态和键值；应用运行时，空格/W/上方向键能够触发跳跃，R/回车能够重启。该测试把物理串行协议、外设寄存器和软件动作三部分串联起来。", body)
    add_text(doc, reference, "5.1.4 中断、异常与返回测试", h2, outline=2)
    add_text(doc, reference, "中断测试重点验证三类 trap 的入口、优先级和返回路径。非法指令异常和 ECALL 属于同步事件，分别对应 SCAUSE=01 和 SCAUSE=02，并跳转至 0x00000300、0x00000320；计时器中断属于异步事件，只有 INTMASK[6] 为 1 且 STATUS[0] 为 0 时才响应，入口为 0x00000340。顶层测试平台可用 CHECK_DINO_TICK 自动检查计时器场景，也可通过 FORCE_INT_START/FORCE_INT_END 对 INT 输入施加单次脉冲。", body)
    add_code(doc, reference, "// top_board_tb.v：中断检查的三个关键观测点\nif (U_TOP.mem_w && (U_TOP.addr_bus == 32'hffff_fe00) &&\n    U_TOP.Cpu_data2bus[0])\n    saw_game_timer_enable = 1;       // 软件已启动游戏定时器\nif (U_TOP.PC == 32'h0000_0340)\n    saw_timer_vector = 1;            // CPU 已跳至计时器入口\nif (U_TOP.mem_w && (U_TOP.addr_bus == 32'he000_0000) &&\n    (U_TOP.Cpu_data2bus == 32'h0000_0001))\n    saw_first_score = 1;              // 返回主循环后得分已推进", body)
    add_text(doc, reference, "对计时器中断，波形中应依次观察到 cpu_timer_irq 脉冲、int_pending[6] 置位、ex_trap_set 与 ex_int_signal 有效、ex_trap_vector=0x00000340、SEPC 和 SCAUSE 更新，以及 STATUS[0] 在返回后清零。图 5-1 和图 5-2 即按这些判据截取。", body)
    add_text(doc, reference, "5.1.5 Dino 应用测试程序", h2, outline=2)
    add_text(doc, reference, "应用测试源文件为 game/game.c，启动和中断入口文件为 game/start.S，构建规则位于 game/Makefile，最终指令初始化文件为 coe/board/I_dino_game.coe。游戏启动时先清空并绘制文本场景，初始化得分和障碍位置，再向 INTMASK 写入 0x40 并使能游戏定时器。主循环持续读取 PS/2 状态，但游戏物理和显示更新只在 frame_ticks 非零时执行，因此键盘事件不会被长延迟阻塞，画面移动又受固定帧节拍约束。", body)
    add_code(doc, reference, "// game/game.c：初始化完成后开启计时器中断\nmmio_write(MMIO_INTMASK, 0x40u);       // 允许第 6 位计时器中断\nmmio_write(MMIO_GAME_TIMER, 1u);      // 启动游戏帧定时器\n\n// 每次中断服务程序发布节拍后，仅更新一帧游戏状态\nif (frame_ticks == 0u)\n    continue;                         // 尚无新帧，继续处理键盘输入\nframe_ticks = 0u;                     // 消费本次帧节拍\nupdate_obstacle_and_score();          // 更新障碍物和得分", body)
    add_text(doc, reference, "应用测试检查 READY、运行、跳跃、障碍移动、碰撞、GAME OVER、重启和得分显示等状态转换。VGA 负责显示角色、地面、障碍和文字；LED 与数码管提供辅助状态。游戏测试通过并不替代 CPU 单元测试，而是作为对指令执行、MMIO、键盘、定时器和中断返回的综合验证。", body)
    add_table(doc, reference, ["测试层次", "测试文件/程序", "输入或刺激", "主要判据"], [
        ("CPU 基础功能", "sim/sccomp_tb.v、课程测试 COE", "复位、指令序列", "写回、分支、访存和阶段标记正确"),
        ("顶层 MMIO", "sim/top_board_tb.v", "COE、拨码、外设模型", "显示/VGA 写入与地址译码正确"),
        ("PS/2", "top_board_tb.v 的 send_ps2_byte", "扫描码 0x29 等", "ready、键值和应用动作一致"),
        ("trap", "run_top_board_sim.py + CHECK_DINO_TICK", "定时器 IRQ 或强制脉冲", "pending、0x340、SEPC、SCAUSE、返回正确"),
        ("Dino 应用", "game/game.c、I_dino_game.coe", "键盘与计时器节拍", "状态、得分、碰撞和重启正确"),
    ], cap, [2.4, 4.2, 3.8, 4.7])
    add_text(doc, reference, "表 5-1 测试代码与结果", cap, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, reference, "5.2 计时器中断测试结果", h2, outline=1)
    add_text(doc, reference, "图 5-1 显示了中断响应的局部时序。cpu_timer_irq 拉高后，int_pending[6] 置位为 40；在 EX 阶段，ex_trap_set 与 ex_int_signal 同时有效，ex_trap_vector 输出 00000340。PC 从正常程序地址转入 00000340，SEPC 保存为 00000704，SCAUSE 被写入 06，STATUS 置为 01。说明处理器完成了计时器中断的请求锁存、向量跳转和现场保存。", body)
    add_picture(doc, reference, zoom, 15.4)
    add_text(doc, reference, "图 5-1 计时器中断响应与向量跳转波形", cap, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, reference, "图 5-2 给出了同一次中断更长时间尺度上的波形。处理程序执行后 STATUS[0] 由 01 回到 00，表明 ERET/ERETN 使 CPU 退出单级 trap 状态并恢复主程序执行。", body)
    add_picture(doc, reference, ret, 15.4)
    add_text(doc, reference, "图 5-2 计时器中断服务程序返回时序", cap, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, reference, "5.3 应用测试结果分析", h2, outline=1)
    add_text(doc, reference, "Dino 程序启动后先初始化文本显存、得分区和地面，再使能 INTMASK 和游戏定时器。键盘输入通过 PS/2 MMIO 读取，空格、W 或上方向键触发跳跃，R 或回车使游戏回到起始状态。每次帧到达后程序更新恐龙高度、仙人掌位置和得分；发生碰撞后显示 GAME OVER 并驱动 LED。顶层仿真日志中 CHECK_DINO_TICK 通过，说明计时器中断已推动首帧和得分变化。", body)

    # 6 Summary.
    add_text(doc, reference, "6 实验总结", h1, outline=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, reference, "6.1 实验中遇到的问题及解决方法", h2, outline=1)
    add_table(doc, reference, ["问题", "解决过程"], [
        ("VGA 扫描时序不稳定", "先验证固定显示图案，再检查行场计数和像素时钟使能，最终按标准 640×480 时序调整扫描模块。"),
        ("PS/2 按键识别不稳定", "将断码与扩展前缀作为独立状态处理，利用扫描记录判断新事件，避免重复读取同一按键。"),
        ("流水线相关导致结果错误", "逐周期检查 ID/EX、EX/MEM、MEM/WB 的寄存器号，补充前推优先级，并对 load-use 情形插入气泡。"),
        ("中断容易与分支控制冲突", "将同步异常、计时器中断、跳转和停顿的优先级集中在 EX 阶段处理；trap 时清空前级并抑制当前指令副作用。"),
        ("游戏画面存在残影", "记录对象的前一帧位置，只擦除旧单元并重绘变化区域，减少整屏 MMIO 写入。"),
    ], cap, [4.4, 10.7])
    add_text(doc, reference, "表 6-1 实验问题与解决方法", cap, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, reference, "6.2 取得的收获", h2, outline=1)
    add_text(doc, reference, "通过本次实验，我对处理器设计的理解从单个模块扩展到完整系统。前推、停顿和冲刷只有放到具体周期和具体指令序列中才能判断是否正确；中断也不仅是 PC 跳转，还涉及 pending、优先级、现场保存、返回地址和对流水线副作用的控制。调试过程中，把问题拆成时钟与显示、MMIO 读写、应用状态和 trap 时序等小范围验证，比同时修改多个模块更有效。小组协作中，PS/2、VGA 和游戏应用的接口约定同样重要，接口稳定后才能进行可靠集成。", body)

    # References: preserve title and teacher-review page already in the template.
    add_page_break(doc, reference)
    for p in list(doc.paragraphs):
        if p.text.startswith("[1] David Patterson"):
            p._element.getparent().remove(p._element)
    for ref in [
        "[1] Patterson D A, Hennessy J L. 计算机组成与设计：硬件/软件接口（RISC-V版）[M]. 第2版. 北京：机械工业出版社，2023.",
        "[2] RISC-V International. The RISC-V Instruction Set Manual, Volume I: Unprivileged ISA [EB/OL]. https://riscv.org/technical/specifications/.",
        "[3] AMD Xilinx. Vivado Design Suite User Guide: Synthesis (UG901) [EB/OL]. 2018.3.",
        "[4] Digilent. Nexys A7 FPGA Trainer Board Reference Manual [EB/OL]. https://digilent.com/reference/programmable-logic/nexys-a7/start.",
        "[5] 计算机系统综合设计课程实验指导材料与课堂课件[Z].",
    ]:
        add_text(doc, reference, ref, body)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()

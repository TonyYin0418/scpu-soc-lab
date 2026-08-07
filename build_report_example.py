from pathlib import Path
from subprocess import run

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "docs" / "雷军班_尹玉文东_2024302131025_实验报告_示例填空稿.docx"
ASSET = Path("/private/tmp") / "scpu_report_architecture.png"


def set_font(run_, name="宋体", size=12, bold=False, color=None):
    run_.font.name = name
    run_._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run_._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run_._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run_.font.size = Pt(size)
    run_.font.bold = bold
    if color:
        run_.font.color.rgb = RGBColor(*color)


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, size=9.5, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_ = paragraph.add_run("第 ")
    set_font(run_, size=9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run_ = paragraph.add_run(" 页")
    set_font(run_, size=9)


def field(paragraph, instruction):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instruction)
    paragraph._p.append(fld)


def add_paragraph(doc, text="", style=None, indent=True, align=None, before=0, after=6, line=1.5):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if indent:
        pf.first_line_indent = Cm(0.74)
    r = p.add_run(text)
    set_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, name="黑体", size={1: 16, 2: 14, 3: 12}[level], bold=True)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_font(r, size=10)
    return p


def add_code(doc, code):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    shade(cell, "F2F2F2")
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(code)
    set_font(r, name="Menlo", size=8.5)
    return table


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        shade(cell, "D9EAF7")
        set_cell_text(cell, h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        if widths:
            cell.width = Cm(widths[j])
    for row in rows:
        cells = table.add_row().cells
        for j, value in enumerate(row):
            set_cell_text(cells[j], value, size=8.7, align=WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT)
            if widths:
                cells[j].width = Cm(widths[j])
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "start", "bottom", "end"):
                node = OxmlElement(f"w:{side}")
                node.set(qn("w:w"), "80")
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            tc_pr.append(margins)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_placeholder(doc, label):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    shade(cell, "FFF2CC")
    set_cell_text(cell, "【待替换图：" + label + "】", bold=True, size=11,
                  color=(156, 87, 0), align=WD_ALIGN_PARAGRAPH.CENTER)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(40)
    p.paragraph_format.space_after = Pt(40)
    return table


def make_architecture():
    dot = r'''digraph G {
      graph [rankdir=LR, bgcolor="white", pad="0.15", nodesep="0.28", ranksep="0.42"];
      node [shape=box, style="rounded,filled", fontname="PingFang SC", fontsize=12, color="#4F81BD", fillcolor="#EAF2F8", margin="0.12,0.08"];
      edge [fontname="PingFang SC", fontsize=9, color="#666666", arrowsize=0.7];
      ps2 [label="PS/2 键盘\nMMIO"];
      timer [label="游戏定时器\nIRQ"];
      cpu [label="五级流水线 RISC-V CPU\nIF / ID / EX / MEM / WB\n冒险处理 + 单级中断/异常", fillcolor="#D9EAD3", color="#6AA84F"];
      imem [label="指令 ROM\nCOE"];
      dmem [label="数据 RAM"];
      mio [label="MIO 总线\n地址译码"];
      vga [label="VGA 文本显示\n80×60"];
      gpio [label="LED / 数码管"];
      game [label="Dino 应用\n状态机与帧更新", fillcolor="#FCE5CD", color="#E69138"];
      imem -> cpu [label="inst"];
      cpu -> dmem [label="地址/读写数据"];
      cpu -> mio [label="MMIO"];
      ps2 -> mio;
      timer -> cpu [label="INT"];
      mio -> vga;
      mio -> gpio;
      cpu -> game [style=dashed, label="执行"];
    }'''
    dot_path = Path("/private/tmp/scpu_report_architecture.dot")
    dot_path.write_text(dot, encoding="utf-8")
    run(["dot", "-Tpng", "-Gdpi=170", str(dot_path), "-o", str(ASSET)], check=True)


def setup_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.7)
    sec.right_margin = Cm(2.7)
    footer = sec.footer
    add_page_number(footer.paragraphs[0])

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    for n, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        s = styles[n]
        s.font.name = "黑体"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        s.font.size = Pt(size)
        s.font.bold = True
    return doc


def new_page(doc):
    """Start a physical Word page reliably across Word and macOS Quick Look."""
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.7)
    sec.right_margin = Cm(2.7)
    return sec


def cover(doc):
    for _ in range(3): doc.add_paragraph()
    for text, size in (("武汉大学计算机学院", 22), ("本科生课程设计报告", 24)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_font(r, name="黑体", size=size, bold=True)
        p.paragraph_format.space_after = Pt(14)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("基于 RISC-V 的 SoC 系统设计")
    set_font(r, name="黑体", size=22, bold=True)
    p.paragraph_format.space_before = Pt(38)
    p.paragraph_format.space_after = Pt(44)
    info = [
        ("专 业 名 称", "【待填：专业名称】"),
        ("课 程 名 称", "计算机系统综合设计"),
        ("指 导 教 师", "蔡朝晖"),
        ("学 生 学 号", "2024302131025"),
        ("学 生 姓 名", "尹玉文东"),
        ("小 组 成 员", "钱瑞恒"),
    ]
    for k, v in info:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{k}：{v}")
        set_font(r, size=14)
        p.paragraph_format.space_after = Pt(10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("二○二六年七月")
    set_font(r, size=14)
    p.paragraph_format.space_before = Pt(36)
    new_page(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("郑 重 声 明")
    set_font(r, name="黑体", size=18, bold=True)
    p.paragraph_format.space_after = Pt(28)
    add_paragraph(doc, "本示例稿用于展示课程设计报告的章节组织、图表位置和写作粒度，不作为验收记录或提交版本。正式提交前，应由小组成员依据实际代码、实机照片、仿真波形和分工情况逐项替换、核对，并在本页按学校要求签字。", before=0, after=12)
    add_paragraph(doc, "正式报告应如实说明小组协作内容、引用资料和测试条件；所有图片、数据与结论均应能追溯到对应的实验过程。", before=0, after=12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("本人签名：______________    日期：______________")
    set_font(r, size=12)
    new_page(doc)


def abstract_and_toc(doc):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("摘  要"), name="黑体", size=18, bold=True)
    add_paragraph(doc, "本项目以 RV32I 指令集为基础，完成了一套面向 Nexys A7 开发板的五级流水线 SoC，并在其上运行文本模式的 Dino 小游戏。处理器由取指、译码、执行、访存和写回五个阶段组成，针对数据相关设置前推与冒险检测逻辑；系统通过 MIO 总线连接指令存储器、数据存储器、PS/2 键盘、VGA 文本显示、LED、数码管以及计时器。示例方案将计时器产生的中断作为帧更新节拍：中断入口保存现场并发布帧到达标志，应用主循环据此完成障碍移动、跳跃状态更新、碰撞判断和得分刷新。", indent=True)
    add_paragraph(doc, "报告重点说明处理器总体结构、地址映射、关键模块实现方法和测试组织方式。由于本文件是填空示例，图 5-1 至图 5-4、实际中断源数量、软件版本号和个人问题记录均需在提交前以真实材料替换。", indent=True)
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    r = p.add_run("关键词：") ; set_font(r, bold=True)
    set_font(p.add_run("RISC-V；五级流水线；SoC；中断；VGA；Nexys A7"))
    new_page(doc)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("目  录"), name="黑体", size=18, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    field(p, 'TOC \\o "1-3" \\h \\z \\u')
    note = doc.add_paragraph(); note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(note.add_run("打开 Word 后右键目录，选择“更新域—更新整个目录”。"), size=10, color=(128, 128, 128))
    new_page(doc)


def introduction(doc):
    add_heading(doc, "1 引言", 1)
    add_heading(doc, "1.1 实验目的", 2)
    add_paragraph(doc, "本实验的目标是将处理器课程中的数据通路、控制器、流水线冒险和异常处理等知识落实到可运行的硬件系统中。除了完成 RV32I 基础指令的执行，还需要建立从交叉编译、COE 初始化到 FPGA 下载和外设验证的完整链路。最终以 Dino 小游戏为应用程序，检验处理器对分支跳转、访存、键盘输入、显示输出以及定时事件的综合处理能力。")
    add_heading(doc, "1.2 国内外研究现状", 2)
    add_paragraph(doc, "RISC-V 采用开放指令集规范，便于高校在不依赖专有指令授权的前提下开展处理器结构、编译器和软硬件协同实验。国外的教学实践通常以单周期处理器为起点，再逐步加入流水线、缓存、特权机制和外设；国内课程中也常使用 FPGA 将 Verilog 设计与可见外设结合，使学生能够观察程序从指令执行到屏幕输出的全过程。")
    add_paragraph(doc, "本项目不追求通用操作系统级功能，而是采用小规模、可综合的 SoC 作为教学实现对象。这样的范围便于把重点放在流水线控制、内存映射 I/O、中断入口和应用程序验证上，也便于根据波形、串行日志或屏幕现象定位问题。")


def environment(doc):
    add_heading(doc, "2 实验环境介绍", 1)
    add_heading(doc, "2.1 Verilog HDL", 2)
    add_paragraph(doc, "Verilog HDL 用于描述处理器数据通路、控制逻辑、外设接口和仿真测试平台。设计以模块化方式组织，顶层负责时钟、复位和外设连接，CPU 内部按流水线阶段传递寄存器和控制信号。")
    add_heading(doc, "2.2 RISC-V 汇编与交叉编译工具", 2)
    add_paragraph(doc, "应用程序使用 RV32I 汇编和 C 语言编写。构建流程将源程序链接为裸机 ELF，再转换为指令 ROM 所需的 COE 文件。正式版本应在此处填写实际使用的交叉编译器名称、版本及生成命令。")
    add_heading(doc, "2.3 Vivado 2018.3", 2)
    add_paragraph(doc, "Vivado 2018.3 用于综合、实现、生成比特流以及下载到开发板。仿真阶段通过测试平台观察 PC、访存控制、VGA 写入和中断入口；实板阶段通过 VGA 显示器、键盘、LED 与数码管核验应用行为。")
    add_heading(doc, "2.4 Nexys A7", 2)
    add_paragraph(doc, "Nexys A7 作为 FPGA 目标平台，提供时钟、按键、拨码开关、LED、七段数码管、VGA 和 PS/2 接口。本项目将 VGA 用于 80×60 文本显示，将 PS/2 用于游戏输入，并通过 LED 和数码管显示辅助状态和得分。")
    add_table(doc, ["类别", "实际环境/待补充项", "用途"], [
        ("硬件平台", "Nexys A7", "下载 SoC，连接 VGA 与 PS/2 键盘"),
        ("FPGA 工具", "Vivado 2018.3", "综合、实现、仿真、生成 bitstream"),
        ("硬件描述", "Verilog HDL", "CPU、外设与顶层设计"),
        ("软件工具", "【待填：交叉编译器、版本】", "编译 RV32I 游戏程序并生成 COE"),
        ("协作成员", "尹玉文东、钱瑞恒", "按实际分工补充"),
    ], [2.7, 7.0, 6.0])
    add_caption(doc, "表 2-1 实验环境")


def architecture(doc):
    add_heading(doc, "3 系统架构设计", 1)
    add_heading(doc, "3.1 总体设计", 2)
    add_paragraph(doc, "示例系统采用五级流水线结构：IF 阶段从指令 ROM 取指，ID 阶段完成译码、寄存器读取和立即数生成，EX 阶段执行算术逻辑、分支判断及异常/中断判定，MEM 阶段访问数据存储器或 MIO 总线，WB 阶段将结果写回寄存器堆。针对 RAW 数据相关，前推单元优先从 EX/MEM 或 MEM/WB 选择最新结果；对 load-use 冒险，冒险检测单元暂停前级并向后级注入气泡。")
    add_paragraph(doc, "在示例中，计时器以固定周期向 CPU 的 INT 输入发出请求。EX 阶段在当前指令无更高优先级同步异常时响应中断，将 PC 转移至计时器向量入口；软件中断处理程序保存现场、更新帧到达标志并执行返回指令。注意：正式报告必须把本段中的中断源、向量地址、返回指令和实际测试结果替换为本组最终实现。")
    make_architecture()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSET), width=Cm(15.4))
    add_caption(doc, "图 3-1 SoC 总体架构（示例）")
    add_heading(doc, "3.2 地址映射与中断向量", 2)
    add_table(doc, ["地址/向量", "模块", "作用"], [
        ("0x0000_0000 起", "指令 ROM", "保存 RV32I 指令及中断入口代码"),
        ("0xC000_0000", "VGA 文本显存", "按单元写入字符和属性"),
        ("0xD000_0000", "PS/2 键盘", "读取扫描码和就绪状态"),
        ("0xE000_0000", "GPIO/七段数码管", "读取按键开关，输出数码管"),
        ("0xF000_0000", "LED", "显示辅助状态或得分低位"),
        ("0xFFFF_FE00", "游戏定时器", "【示例】使能/关闭帧节拍"),
        ("0xFFFF_FF00", "INTMASK", "【示例】设置计时中断使能位"),
        ("0x0000_0300", "非法指令入口", "【示例】同步异常入口"),
        ("0x0000_0320", "ECALL 入口", "【示例】系统调用异常入口"),
        ("0x0000_0340", "计时器入口", "【示例】异步计时中断入口"),
    ], [3.1, 4.1, 8.5])
    add_caption(doc, "表 3-1 内存映射与向量入口（示例，需按最终实现复核）")
    add_heading(doc, "3.3 功能模块及接口", 2)
    add_paragraph(doc, "以下表格覆盖总体架构图中的主要功能单元。端口仅列出对理解模块作用最重要的信号，正式报告可按最终 RTL 的端口名补全或修订。")
    add_table(doc, ["模块", "功能描述", "主要输入", "主要输出"], [
        ("PC/NPC", "保存当前指令地址，选择顺序、分支、跳转或 trap 目标。", "clk、reset、redirect、trap_vector", "PC_out、PC+4"),
        ("指令 ROM", "按 PC 提供 32 位指令。", "PC[11:2]", "inst"),
        ("控制器与立即数扩展", "识别操作码，生成 ALU、访存、写回和分支控制信号。", "inst", "ALUOp、RegWrite、MemWrite、imm"),
        ("RF", "保存 32 个通用寄存器，x0 恒为 0。", "rs1、rs2、rd、wd、RegWrite", "rs1_data、rs2_data"),
        ("ALU", "完成算术、逻辑、比较和地址计算。", "ALU_A、ALU_B、ALUOp", "alu_result、zero"),
        ("前推单元", "解决可前推的 RAW 相关，选择最新操作数。", "ID/EX、EX/MEM、MEM/WB 寄存器号", "forwardA、forwardB"),
        ("冒险检测", "检测 load-use 等不能前推的相关并插入气泡。", "ID rs、EX rd、MemRead", "stall、flush"),
        ("DM/MIO 控制", "完成字节/半字/字访问，并分发到 RAM 或外设。", "addr、data、mem_w、dm_ctrl", "Data_read、外设写使能"),
        ("异常/中断控制", "保存 trap 状态，选择向量并冲刷流水线。", "EX cause、INT、INTMASK、STATUS", "trap_vector、flush、SEPC"),
        ("PS/2 接口", "接收扫描码，为软件提供键盘事件。", "PS2C、PS2D、read", "key、ready、scan_log"),
        ("VGA 文本显示", "把字符单元转换为像素时序和 RGB 输出。", "MMIO 写入、像素时钟", "VGA_R/G/B、HS、VS"),
        ("游戏定时器", "【示例】生成固定频率的帧中断。", "clk、enable", "tick_irq"),
    ], [2.4, 5.2, 4.2, 4.1])
    add_caption(doc, "表 3-2 功能模块与主要接口")


def key_modules(doc):
    add_heading(doc, "4 关键模块设计", 1)
    add_heading(doc, "4.1 DM 读写控制器", 2)
    add_paragraph(doc, "DM 读写控制器根据访问类型产生字节使能、地址低位选择和读数据扩展方式。对 lb/lbu、lh/lhu 等读指令，需要从 32 位返回数据中截取目标字节或半字，再按有符号或无符号方式扩展；对 sb/sh/sw 等写指令，需要将写数据对齐到相应的字节通道。这样可以保证小端存储下不同宽度的访存语义一致。")
    add_code(doc, "// 示例：按地址低两位选择一个字节，并按符号位扩展\ncase (addr[1:0])\n  2'b00: byte_data = Data_in[7:0];\n  2'b01: byte_data = Data_in[15:8];\n  2'b10: byte_data = Data_in[23:16];\n  default: byte_data = Data_in[31:24];\nendcase\nread_data = unsigned_load ? {24'b0, byte_data} : {{24{byte_data[7]}}, byte_data};")
    add_heading(doc, "4.2 前推单元", 2)
    add_paragraph(doc, "当 EX 阶段源寄存器与后续流水寄存器中的目的寄存器相同时，若结果已经在 EX/MEM 或 MEM/WB 阶段产生，则无需等待写回寄存器堆。前推单元优先选择更靠近 EX 阶段的 EX/MEM 结果，再选择 MEM/WB 写回数据；目的寄存器为 x0 时不参与前推。")
    add_code(doc, "// 示例：EX/MEM 优先级高于 MEM/WB\nif (ex_mem_regwrite && ex_mem_rd != 0 && ex_mem_rd == id_ex_rs1)\n    forwardA = 2'b10;\nelse if (mem_wb_regwrite && mem_wb_rd != 0 && mem_wb_rd == id_ex_rs1)\n    forwardA = 2'b01;\nelse\n    forwardA = 2'b00;")
    add_heading(doc, "4.3 冒险检测", 2)
    add_paragraph(doc, "load 指令的数据直到 MEM 阶段末才可用，即使设置前推，紧随其后的使用者也无法在 EX 阶段取得正确操作数。冒险检测单元在 ID 阶段发现该类 load-use 相关后保持 PC 和 IF/ID 寄存器不变，同时向 ID/EX 写入气泡。下一周期数据进入可前推位置后，相关指令即可继续执行。")
    add_code(doc, "// 示例：检测 load-use 冒险\nload_use = id_ex_memread && (id_ex_rd != 0) &&\n           ((id_ex_rd == if_id_rs1) || (id_ex_rd == if_id_rs2));\nif (load_use) begin\n    pc_write = 1'b0;\n    if_id_write = 1'b0;\n    id_ex_flush = 1'b1;\nend")
    add_heading(doc, "4.4 单级中断与异常控制器（示例）", 2)
    add_paragraph(doc, "示例控制器在 EX 阶段统一判定 trap。非法指令和 ECALL 属于同步异常，优先于异步计时中断；中断允许位由 INTMASK 控制。发生 trap 时保存当前 EX 指令 PC 至 SEPC，记录 SCAUSE，置位 STATUS.EXL，并清空 IF/ID、ID/EX 以及当前 EX 指令的提交结果，避免错误路径产生寄存器或内存副作用。处理程序完成后由课程定义的返回指令恢复 PC。")
    add_code(doc, "// 示例：同步异常优先于计时中断\nif (!STATUS[0]) begin\n  if (EX_SCAUSE != SCAUSE_NONE) begin\n      trap_vector = exception_vector; trap_set = 1'b1;\n  end else if (INT_PEND[6] && INTMASK[6]) begin\n      trap_vector = 32'h0000_0340; trap_set = 1'b1;\n      int_signal = 1'b1;\n  end\nend")
    add_heading(doc, "4.5 VGA 与 Dino 帧更新（示例）", 2)
    add_paragraph(doc, "VGA 模块把文本显存中的字符和属性编码转换为像素输出。应用层把屏幕看作 80×60 个单元，单元地址按 row×80+col 计算。Dino 状态机包含 READY、RUNNING 和 GAME OVER 等状态；收到帧到达标志后更新恐龙高度、障碍横坐标、得分和碰撞状态，仅重绘发生变化的区域。")
    add_code(doc, "// 示例：中断只发布节拍，主循环完成一帧更新\ntimer_trap: save_registers(); frame_ticks = frame_ticks + 1; restore_registers(); eret();\n\nif (frame_ticks != 0) begin\n    frame_ticks = 0;\n    update_physics(); update_obstacle();\n    render_changed_cells();\nend")


def testing(doc):
    add_heading(doc, "5 测试及结果分析", 1)
    add_heading(doc, "5.1 测试代码及分析", 2)
    add_paragraph(doc, "处理器基础功能可由课程统一测试程序和自建仿真平台验证；应用部分使用 Dino 程序作为综合测试。应保留测试文件名、构建命令、COE 文件名和必要的输入条件，而不必粘贴整个交叉编译后的汇编。以下清单为示例，正式提交前必须按实际工程文件和执行记录修改。")
    add_table(doc, ["测试项", "示例测试文件/方式", "观察内容"], [
        ("流水线基础指令", "sim/sccomp_tb.v；课程统一测试 COE", "寄存器写回、分支、访存结果"),
        ("数据冒险", "自建相关指令序列", "前推结果、load-use 停顿周期"),
        ("中断/异常", "【待填：testbench 或 COE 文件名】", "INT、PC 跳转、SEPC/SCAUSE、返回后续执行"),
        ("PS/2 输入", "top_board_tb.v + 键盘扫描码刺激", "空格/W/方向键对应的游戏动作"),
        ("VGA 与游戏", "game/game.c；I_dino_game.coe", "启动画面、跳跃、障碍、得分、结束和重启"),
    ], [3.0, 6.0, 6.5])
    add_caption(doc, "表 5-1 测试项目与观察点（示例）")
    add_heading(doc, "5.2 中断路径测试（示例写法）", 2)
    add_paragraph(doc, "【待替换为实际测试结论】向计时器使能寄存器写入使能位后，定时器开始产生周期性 INT 请求。波形中应能观察到 INT 被锁存为 pending，CPU 在 EX 阶段完成当前控制决策后将 PC 转向 0x00000340；处理程序更新 frame_ticks 并执行返回指令，随后主程序继续运行。若同时存在非法指令或 ECALL，应以同步异常优先，分别跳转至 0x00000300 和 0x00000320。")
    add_placeholder(doc, "中断仿真波形：INT、PC、SEPC、SCAUSE、frame_ticks 与返回路径")
    add_caption(doc, "图 5-1 中断响应与返回波形（待插入真实截图）")
    add_heading(doc, "5.3 Dino 应用测试结果（示例写法）", 2)
    add_paragraph(doc, "【待替换为实际测试结论】系统上电后，VGA 显示 Dino Runner 起始界面和初始得分；按下跳跃键后，恐龙位置随帧更新改变，仙人掌从右向左移动，得分同步写入 VGA 和数码管。发生碰撞时显示 GAME OVER 并点亮 LED；执行重启输入后，状态、得分和障碍位置恢复为初值。每张图片都应在图注中说明输入条件、显示现象和对应的功能结论。")
    add_placeholder(doc, "Nexys A7 实机运行照片：READY/开始界面")
    add_caption(doc, "图 5-2 Dino 启动界面（待插入真实照片）")
    add_placeholder(doc, "Nexys A7 实机运行照片：跳跃过程与得分变化")
    add_caption(doc, "图 5-3 Dino 运行过程（待插入真实照片）")
    add_placeholder(doc, "Nexys A7 实机运行照片：碰撞后的 GAME OVER 界面")
    add_caption(doc, "图 5-4 Dino 碰撞结果（待插入真实照片）")


def summary(doc):
    add_heading(doc, "6 实验总结", 1)
    add_heading(doc, "6.1 实验中遇到的问题及解决方法", 2)
    add_paragraph(doc, "本节必须由实际参与成员根据真实过程改写。下面保留的是根据版本记录整理的示例条目，不应直接作为个人经历提交。")
    add_table(doc, ["问题现象", "排查与原因", "解决方法"], [
        ("VGA 画面不稳定或显示区域不正确", "核对像素时钟、行场计数和有效显示区，发现扫描逻辑与标准 640×480 时序不一致。", "调整 VGA 扫描时序，并使用时钟使能控制扫描更新。"),
        ("键盘按键在软件中识别异常", "PS/2 数据包含断码和扩展前缀，若只读取最后一个字节会误判。", "在软件或接口层保存前缀状态，区分按下、释放和扩展按键。"),
        ("流水线相关指令结果错误", "追踪 ID/EX、EX/MEM、MEM/WB 的寄存器号，发现未优先使用最新结果或遗漏 load-use 情形。", "补充前推优先级，并对 load-use 相关插入一个气泡。"),
        ("更新游戏画面后出现残影或闪烁", "整屏清空会带来大量 MMIO 写入，且容易覆盖尚需保留的元素。", "记录对象上一帧位置，仅擦除旧单元并重绘变化区域。"),
        ("【待填：你本人遇到的问题】", "【待填：如何定位】", "【待填：具体修复和验证方式】"),
    ], [4.2, 6.2, 5.1])
    add_caption(doc, "表 6-1 问题与解决方法（示例，需按真实经历替换）")
    add_heading(doc, "6.2 取得的收获", 2)
    add_paragraph(doc, "【需由本人改写】本次实验把原先分散的处理器知识串成了完整流程：从 Verilog 模块、仿真波形到 FPGA 外设，再到交叉编译的应用程序。实际调试时，最有帮助的不是一次性改动很多模块，而是先缩小问题范围，例如先验证时钟和 VGA 时序，再验证 MMIO 读写，最后观察游戏状态变化。对流水线而言，前推、停顿和冲刷必须结合具体周期来理解；对中断而言，入口、现场保护、返回地址和流水线副作用需要一起检查。")
    add_paragraph(doc, "【待补充：个人真实收获、承担工作以及仍希望改进的部分。建议写 150—250 字，不使用空泛套话。】")


def references_and_review(doc):
    add_heading(doc, "参考文献", 1)
    refs = [
        "[1] Patterson D A, Hennessy J L. 计算机组成与设计：硬件/软件接口（RISC-V 版）[M]. 第2版. 北京：机械工业出版社，2023.",
        "[2] RISC-V International. The RISC-V Instruction Set Manual, Volume I: Unprivileged ISA [EB/OL]. https://riscv.org/technical/specifications/.",
        "[3] AMD Xilinx. Vivado Design Suite User Guide: Synthesis (UG901) [EB/OL]. 2018.3.",
        "[4] Digilent. Nexys A7 FPGA Trainer Board Reference Manual [EB/OL]. https://digilent.com/reference/programmable-logic/nexys-a7/start.",
        "[5] 课程实验指导材料与课堂课件（请按实际资料名称、版本和访问日期补充）。",
    ]
    for ref in refs:
        add_paragraph(doc, ref, indent=False, after=6, line=1.25)
    new_page(doc)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("教师评语评分"), name="黑体", size=18, bold=True)
    for _ in range(3): doc.add_paragraph()
    p = doc.add_paragraph(); set_font(p.add_run("评语："), size=12)
    for _ in range(12): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(p.add_run("评分：______________"), size=12)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(p.add_run("评阅人：______________"), size=12)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(p.add_run("年    月    日"), size=12)


def main():
    doc = setup_doc()
    cover(doc)
    abstract_and_toc(doc)
    introduction(doc)
    environment(doc)
    architecture(doc)
    key_modules(doc)
    testing(doc)
    summary(doc)
    references_and_review(doc)
    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
